"""Shared geometry, colors, and shape helpers for build_pptx.py.

Mirrors deck/style.css so the .pptx reads as the same deck, translated into
native (fully editable) PowerPoint text boxes, tables, shapes, and charts.

PPTX has no flow layout, so every block below is placed at an explicit y via
a small `Cursor` (a mutable running y-offset per column) and advances it by
an *estimated* rendered height (word-wrap has no offline layout engine to
query, so heights are approximated from character count / column width).
Good enough to land close; the point of the export is to finish by hand in
PowerPoint, not to be pixel-perfect out of the box.
"""
from __future__ import annotations
import math
import re
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ colors
def C(hexstr):
    return RGBColor.from_string(hexstr)

INK = C("0B0B0B")
INK2 = C("52514E")
MUTED = C("898781")
GRID = C("E1E0D9")
RULE = C("C3C2B7")
SURF = C("FCFCFB")
BODY_BG = C("F2F2EF")
ACCENT = C("2A78D6")
CRIT = C("D03B3B")
GOOD = C("0CA30C")
ORANGE = C("EB6834")
DEEMPH = C("D9D8D2")
WASH = C("CDE2FB")
TOTAL_GREY = C("6B6A66")

PILL = {
    "pass": (C("E4F5E4"), C("0A6B0A")),
    "fix": (C("E9F1FD"), C("1C5CAB")),
    "warn": (C("FDF0E6"), C("9C4A18")),
}

FONT = "Calibri"

# ------------------------------------------------------------------ geometry
SLIDE_W = 13.333
SLIDE_H = 7.5
PAD_L = 0.55
PAD_R = 0.55
PAD_T = 0.35
PAD_B = 0.28
CONTENT_W = SLIDE_W - PAD_L - PAD_R
GAP = 0.24

EYEBROW_Y = PAD_T
EYEBROW_H = 0.16
TITLE_Y = EYEBROW_Y + EYEBROW_H + 0.02
TITLE_H_1LINE = 0.32
FOOT_H = 0.22


def col_geometry(widths_pct, body_y, body_bottom):
    n = len(widths_pct)
    total_gap = GAP * (n - 1)
    usable = CONTENT_W - total_gap
    xs = []
    x = PAD_L
    for pct in widths_pct:
        w = usable * (pct / 100.0)
        xs.append((x, w))
        x += w + GAP
    h = body_bottom - body_y
    return [(x, body_y, w, h) for x, w in xs]


class Cursor:
    def __init__(self, x, y, w):
        self.x, self.y, self.w = x, y, w

    def advance(self, dy):
        self.y += dy


# ------------------------------------------------------------------ text measurement (heuristic)
TAG_RE = re.compile(r"<[^>]+>")
BR_RE = re.compile(r"<br\s*/?>", re.I)
ENTS = [("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&middot;", "·"),
        ("&rarr;", "→"), ("&mdash;", "—"), ("&ndash;", "–")]


def _plain(seg):
    for a, b in ENTS:
        seg = seg.replace(a, b)
    return TAG_RE.sub("", seg)


def est_lines(html, width_in, size_pt, char_w=0.50):
    chars_per_line = max(6, width_in * 72 / (size_pt * char_w))
    total = 0
    for seg in BR_RE.split(html.replace("\n", "<br>")):
        plain = _plain(seg)
        n = max(1, math.ceil(len(plain) / chars_per_line)) if plain.strip() else 1
        total += n
    return total


def est_text_height(html, width_in, size_pt, line_factor=1.30, char_w=0.50):
    lines = est_lines(html, width_in, size_pt, char_w)
    return lines * (size_pt * line_factor / 72)


# ------------------------------------------------------------------ run parsing
TAG_SPLIT = re.compile(r"(<[^>]+>)")


def _color_for_class(cls):
    if "crit" in cls:
        return CRIT
    if "accent" in cls:
        return ACCENT
    if "good" in cls:
        return GOOD
    if "mute" in cls:
        return MUTED
    return None


def add_runs(paragraph, html, base_size=11, base_color=None, base_bold=False):
    bold = base_bold
    color = base_color
    color_stack = []
    italic = False
    for part in TAG_SPLIT.split(html):
        if not part:
            continue
        if part.startswith("<"):
            inner = part.strip("<>/")
            tag = inner.split()[0].lower() if inner else ""
            closing = part.startswith("</")
            if tag in ("strong", "b"):
                bold = not closing
            elif tag in ("em", "i"):
                italic = not closing
            elif tag == "br":
                paragraph.add_line_break()
            elif tag == "span":
                if closing:
                    color = color_stack.pop() if color_stack else base_color
                else:
                    m = re.search(r'class="([^"]*)"', part)
                    cls = m.group(1) if m else ""
                    color_stack.append(color)
                    c = _color_for_class(cls)
                    color = c if c else color
            continue
        text = _plain(part)
        if not text:
            continue
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(base_size)
        run.font.name = FONT
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color if color else (base_color if base_color else INK2)


def set_no_autofit(tf):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE


def add_textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.1)), Inches(max(h, 0.1)))
    tf = box.text_frame
    set_no_autofit(tf)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


# ------------------------------------------------------------------ page chrome
def add_eyebrow(slide, tag_text, apx=False):
    box, tf = add_textbox(slide, PAD_L, EYEBROW_Y, CONTENT_W, EYEBROW_H)
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = tag_text
    r1.font.size = Pt(8)
    r1.font.bold = True
    r1.font.name = FONT
    r1.font.color.rgb = MUTED if apx else ACCENT
    r2 = p.add_run()
    r2.text = "   Manager, Workforce Management · Care Experience & Strategy"
    r2.font.size = Pt(8)
    r2.font.name = FONT
    r2.font.color.rgb = MUTED
    return box


def add_title(slide, title, apx=False):
    size = 17 if apx else 19
    box, tf = add_textbox(slide, PAD_L, TITLE_Y, CONTENT_W, TITLE_H_1LINE + 0.1)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = INK
    return box


def add_dek(slide, y, dek):
    box, tf = add_textbox(slide, PAD_L, y, CONTENT_W, 0.26)
    p = tf.paragraphs[0]
    add_runs(p, dek, base_size=10.5, base_color=INK2)
    return box


def add_hr(slide, y):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(PAD_L), Inches(y),
                                     Inches(PAD_L + CONTENT_W), Inches(y))
    ln.line.color.rgb = GRID
    ln.line.width = Pt(0.75)
    return ln


def add_footer(slide, left_text, page_no):
    y = SLIDE_H - PAD_B - FOOT_H + 0.05
    add_hr(slide, y - 0.05)
    box, tf = add_textbox(slide, PAD_L, y, CONTENT_W * 0.82, FOOT_H)
    p = tf.paragraphs[0]
    add_runs(p, left_text, base_size=7.5, base_color=MUTED)
    box2, tf2 = add_textbox(slide, PAD_L + CONTENT_W * 0.82, y, CONTENT_W * 0.18, FOOT_H)
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r = p2.add_run()
    r.text = page_no
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = INK2


# ------------------------------------------------------------------ cursor-based content blocks
def h2(slide, cur, text, gap_before=0.06, gap_after=0.05, first=False):
    cur.advance(0 if first else gap_before)
    h = max(0.16, est_text_height(text.upper(), cur.w, 8.2, char_w=0.58))
    box, tf = add_textbox(slide, cur.x, cur.y, cur.w, h)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text.upper()
    r.font.size = Pt(8.2)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = MUTED
    try:
        r._r.get_or_add_rPr().set("spc", "60")
    except Exception:
        pass
    cur.advance(h + gap_after)
    return box


def h3(slide, cur, html, size=11.5, gap_before=0.06, gap_after=0.03):
    cur.advance(gap_before)
    h = est_text_height(html, cur.w, size)
    box, tf = add_textbox(slide, cur.x, cur.y, cur.w, h)
    p = tf.paragraphs[0]
    add_runs(p, html, base_size=size, base_color=INK, base_bold=True)
    cur.advance(h + gap_after)
    return box


def p(slide, cur, html, size=9.5, gap_before=0.05, gap_after=0.0, color=None):
    cur.advance(gap_before)
    h = est_text_height(html, cur.w, size)
    box, tf = add_textbox(slide, cur.x, cur.y, cur.w, h)
    pa = tf.paragraphs[0]
    add_runs(pa, html, base_size=size, base_color=color or INK2)
    cur.advance(h + gap_after)
    return box


def bullets(slide, cur, items, size=9.2, gap_before=0.05, item_gap=0.03):
    cur.advance(gap_before)
    total_h = 0
    box, tf = add_textbox(slide, cur.x, cur.y, cur.w, 0.1)
    first = True
    for it in items:
        seg_h = est_text_height("•  " + it, cur.w, size)
        pa = tf.paragraphs[0] if first else tf.add_paragraph()
        if not first:
            pa.space_before = Pt(item_gap * 72)
        add_runs(pa, "•  " + it, base_size=size, base_color=INK2)
        total_h += seg_h + (item_gap if not first else 0)
        first = False
    box.height = Inches(max(total_h, 0.1))
    cur.advance(total_h + 0.02)
    return box


def hero(slide, cur, number, unit_html, num_size=48, gap_before=0.02, gap_after=0.06):
    cur.advance(gap_before)
    h = num_size * 1.35 / 72 + est_text_height(unit_html, cur.w, 9.5)
    box, tf = add_textbox(slide, cur.x, cur.y, cur.w, h)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    pa = tf.paragraphs[0]
    r = pa.add_run()
    r.text = number
    r.font.size = Pt(num_size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = INK
    r2 = pa.add_run()
    r2.text = "   "
    r2.font.size = Pt(9)
    add_runs(pa, unit_html, base_size=9.5, base_color=INK2, base_bold=True)
    cur.advance(h + gap_after)
    return box


def tiles_row(slide, cur, items, gap_before=0.05, row_h=0.5, gap_after=0.05):
    """items: list of (k, v, s), laid out in a row of equal-width tiles."""
    cur.advance(gap_before)
    n = len(items)
    gap = 0.12
    tw = (cur.w - gap * (n - 1)) / n
    x = cur.x
    for k, v, s in items:
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(cur.y),
                                         Inches(x + tw), Inches(cur.y))
        ln.line.color.rgb = GRID
        ln.line.width = Pt(0.75)
        box, tf = add_textbox(slide, x, cur.y + 0.05, tw, row_h - 0.05)
        pa = tf.paragraphs[0]
        r = pa.add_run()
        r.text = k.upper()
        r.font.size = Pt(6.6)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = MUTED
        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        r2 = p2.add_run()
        r2.text = v
        r2.font.size = Pt(15)
        r2.font.bold = True
        r2.font.name = FONT
        r2.font.color.rgb = INK
        p3 = tf.add_paragraph()
        p3.space_before = Pt(1)
        r3 = p3.add_run()
        r3.text = s
        r3.font.size = Pt(7)
        r3.font.name = FONT
        r3.font.color.rgb = MUTED
        x += tw + gap
    cur.advance(row_h + gap_after)


NOTE_BORDER = {"default": ACCENT, "red": CRIT, "grey": RULE}


def note_open(cur, indent=0.13, gap_before=0.08):
    """For notes with mixed content (heading + paragraph + table, etc.) Call,
    build content on the returned (y0) cursor state, then note_close()."""
    cur.advance(gap_before)
    y0 = cur.y
    cur.x += indent
    cur.w -= indent
    return y0


def note_close(slide, cur, y0, kind="default", gap_after=0.06, indent=0.13):
    cur.x -= indent
    cur.w += indent
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cur.x), Inches(y0),
                                     Inches(cur.x), Inches(cur.y))
    ln.line.color.rgb = NOTE_BORDER.get(kind, ACCENT)
    ln.line.width = Pt(2.25)
    cur.advance(gap_after)


def note(slide, cur, heading_html=None, body_items=None, kind="default",
         h_size=10.5, p_size=8.7, gap_before=0.08, gap_after=0.06, pad=0.03):
    """body_items: list of html paragraphs (each own <p>)."""
    cur.advance(gap_before)
    inner_x = cur.x + 0.13
    inner_w = cur.w - 0.13
    y0 = cur.y
    h_total = pad
    if heading_html:
        h_total += est_text_height(heading_html, inner_w, h_size) + 0.03
    if body_items:
        for it in body_items:
            h_total += est_text_height(it, inner_w, p_size) + 0.02
    h_total += pad
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cur.x), Inches(y0),
                                     Inches(cur.x), Inches(y0 + h_total))
    ln.line.color.rgb = NOTE_BORDER.get(kind, ACCENT)
    ln.line.width = Pt(2.25)
    yy = y0 + pad
    if heading_html:
        hh = est_text_height(heading_html, inner_w, h_size)
        box, tf = add_textbox(slide, inner_x, yy, inner_w, hh)
        pa = tf.paragraphs[0]
        add_runs(pa, heading_html, base_size=h_size, base_color=INK, base_bold=True)
        yy += hh + 0.03
    if body_items:
        for it in body_items:
            hh = est_text_height(it, inner_w, p_size)
            box, tf = add_textbox(slide, inner_x, yy, inner_w, hh)
            pa = tf.paragraphs[0]
            add_runs(pa, it, base_size=p_size, base_color=INK2)
            yy += hh + 0.02
    cur.y = y0 + h_total
    cur.advance(gap_after)


def table(slide, cur, header, rows, col_widths=None, font_size=8.3, header_size=6.8,
          row_h=0.185, total_row_idx=None, gap_before=0.05, gap_after=0.05,
          row_pad=0.055):
    """header: list of (text, align 't'|'n'). rows: list of list of
    (text, align, *modifiers) where modifiers in {'b','crit','good','mute'}.
    row_h is a *floor*; actual row height auto-grows to fit wrapped text so
    rows never overlap the next block."""
    cur.advance(gap_before)
    n_rows = len(rows) + 1
    n_cols = len(header)
    if not col_widths:
        col_widths = [1] * n_cols
    total_cw = sum(col_widths)
    col_w_in = [cur.w * cw / total_cw for cw in col_widths]

    def row_height(cells, size):
        best = row_h
        for j, spec in enumerate(cells):
            text = spec[0] if isinstance(spec, tuple) else spec
            cw = max(col_w_in[j] - 0.09, 0.15)
            lines = est_lines(text.upper() if size == header_size else text, cw, size, char_w=0.56)
            h = lines * (size * 1.28 / 72) + row_pad
            best = max(best, h)
        return best

    heights = [row_height(header, header_size)]
    for row in rows:
        heights.append(row_height(row, font_size))
    h = sum(heights)

    gframe = slide.shapes.add_table(n_rows, n_cols, Inches(cur.x), Inches(cur.y),
                                     Inches(cur.w), Inches(h))
    tbl = gframe.table
    for i, hh in enumerate(heights):
        tbl.rows[i].height = Inches(hh)
    if col_widths:
        acc = 0
        for i, cw in enumerate(col_widths):
            width = int(Inches(cur.w) * cw / total_cw)
            if i == len(col_widths) - 1:
                width = int(Inches(cur.w)) - acc
            tbl.columns[i].width = Emu(width)
            acc += width
    for j, (htext, align) in enumerate(header):
        cell = tbl.cell(0, j)
        cell.margin_left = cell.margin_right = Pt(2.5)
        cell.margin_top = cell.margin_bottom = Pt(1)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURF
        cell.vertical_anchor = MSO_ANCHOR.BOTTOM
        tf = cell.text_frame
        tf.word_wrap = True
        pa = tf.paragraphs[0]
        pa.alignment = PP_ALIGN.RIGHT if align == "n" else PP_ALIGN.LEFT
        r = pa.add_run()
        r.text = htext.upper()
        r.font.size = Pt(header_size)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = MUTED
    for i, row in enumerate(rows):
        is_tot = total_row_idx is not None and i == total_row_idx
        for j, cell_spec in enumerate(row):
            text, align, *rest = cell_spec
            bold = "b" in rest
            pill_mod = next((m for m in rest if m.startswith("pill_")), None)
            color = (PILL[pill_mod[5:]][1] if pill_mod else
                     CRIT if "crit" in rest else GOOD if "good" in rest else
                     MUTED if "mute" in rest else (INK if is_tot else INK2))
            bold = bold or bool(pill_mod)
            cell = tbl.cell(i + 1, j)
            cell.margin_left = cell.margin_right = Pt(2.5)
            cell.margin_top = cell.margin_bottom = Pt(1)
            cell.fill.solid()
            cell.fill.fore_color.rgb = SURF
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            pa = tf.paragraphs[0]
            pa.alignment = PP_ALIGN.RIGHT if align == "n" else PP_ALIGN.LEFT
            r = pa.add_run()
            r.text = text
            r.font.size = Pt(font_size)
            r.font.bold = bold or is_tot
            r.font.name = FONT
            r.font.color.rgb = color
    tblPr = tbl._tbl.find(qn('a:tblPr'))
    if tblPr is not None:
        tblPr.set('firstRow', '0')
        tblPr.set('bandRow', '0')
    cur.advance(h + gap_after)
    return gframe


def pill(slide, x, y, text, kind, w=0.46, h=0.15):
    bg, fg = PILL[kind]
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.5
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Pt(1)
    tf.margin_top = tf.margin_bottom = Pt(0)
    pa = tf.paragraphs[0]
    pa.alignment = PP_ALIGN.CENTER
    r = pa.add_run()
    r.text = text
    r.font.size = Pt(6.3)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = fg
    return box


def image(slide, cur, path, aspect_w, aspect_h, width_frac=1.0, gap_before=0.05, gap_after=0.05, center=True):
    cur.advance(gap_before)
    w = cur.w * width_frac
    h = w * aspect_h / aspect_w
    x = cur.x + (cur.w - w) / 2 if center else cur.x
    slide.shapes.add_picture(path, Inches(x), Inches(cur.y), Inches(w), Inches(h))
    cur.advance(h + gap_after)
    return h
