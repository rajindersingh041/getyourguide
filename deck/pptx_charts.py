"""Native PowerPoint chart builders for build_pptx.py — real chart objects
(editable data, not pictures), styled to match deck/style.css as closely as
python-pptx / OOXML chart formatting allows.
"""
from __future__ import annotations
from pptx.util import Inches, Pt, Emu
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION,
                              XL_MARKER_STYLE, XL_LABEL_POSITION)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

import pptx_helpers as H


def _no_gridlines(chart):
    try:
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.has_minor_gridlines = False
    except Exception:
        pass


def _style_cat_axis(chart, size=7.5, color=H.MUTED):
    ax = chart.category_axis
    ax.tick_labels.font.size = Pt(size)
    ax.tick_labels.font.name = H.FONT
    ax.tick_labels.font.color.rgb = color
    ax.format.line.color.rgb = H.RULE
    ax.major_tick_mark = XL_TICK_LABEL_POSITION.NONE if False else ax.major_tick_mark


def _style_val_axis_hidden(chart):
    chart.value_axis.visible = False


# ------------------------------------------------------------------ waterfall
def waterfall(slide, x, y, w, h, labels, values, kinds, colors, subs=None,
              value_fmt="{:,.0f}", font_size=9):
    """kinds[i] in {'total','delta'}. values are the *magnitude* of each bar
    (delta size, or full height for totals). colors[i] is the fill for bar i."""
    running = 0.0
    base, size = [], []
    for kind, v in zip(kinds, values):
        if kind == "total":
            base.append(0)
            size.append(v)
            running = v
        else:
            base.append(running)
            size.append(v)
            running += v

    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series("base", base)
    cd.add_series("size", size)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y),
                                     Inches(w), Inches(h), cd)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 35
    plot.has_data_labels = False

    s_base, s_size = chart.series[0], chart.series[1]
    s_base.format.fill.background()
    s_base.format.line.fill.background()
    s_size.format.line.fill.background()
    s_size.has_data_labels = True
    s_size.data_labels.font.size = Pt(font_size + 1)
    s_size.data_labels.font.bold = True
    s_size.data_labels.font.name = H.FONT
    s_size.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    # per-point fill + label text on the visible series
    for i, (kind, v, pt) in enumerate(zip(kinds, values, s_size.points)):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i]
        text = value_fmt.format(v) if kind == "total" else "+" + value_fmt.format(v)
        pt.data_label.text_frame.text = text
        pt.data_label.font.size = Pt(font_size + 1)
        pt.data_label.font.bold = True
        pt.data_label.font.name = H.FONT
        pt.data_label.font.color.rgb = H.INK if kind == "total" else H.INK2

    _style_val_axis_hidden(chart)
    _no_gridlines(chart)
    _style_cat_axis(chart, size=font_size - 1.5)
    chart.category_axis.format.line.color.rgb = H.RULE
    return gframe


# ------------------------------------------------------------------ actual+forecast line
def line_with_forecast(slide, x, y, w, h, categories, actual, forecast_incl_join,
                        annotations=None, font_size=7.5):
    """actual: values for the 'actual' months, None elsewhere (aligned to
    categories). forecast_incl_join: values for forecast months *including*
    the last actual month (so the dashed segment visually connects), None
    elsewhere."""
    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("Actual", actual)
    cd.add_series("Forecast", forecast_incl_join)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y),
                                     Inches(w), Inches(h), cd)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    _no_gridlines(chart)
    _style_cat_axis(chart, size=font_size)
    _style_val_axis_hidden(chart)

    s_act, s_fc = chart.series
    for s, dashed in ((s_act, False), (s_fc, True)):
        s.format.line.color.rgb = H.ACCENT
        s.format.line.width = Pt(1.75)
        if dashed:
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        s.marker.style = XL_MARKER_STYLE.CIRCLE
        s.marker.size = 5
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = H.ACCENT
        s.marker.format.line.color.rgb = H.SURF
    return gframe


# ------------------------------------------------------------------ horizontal bar
def hbar(slide, x, y, w, h, categories, values, colors, value_fmt="{:.0f}%",
         font_size=8, cat_font_size=8):
    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("v", values)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y),
                                     Inches(w), Inches(h), cd)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 45
    series = chart.series[0]
    series.has_data_labels = True
    series.data_labels.font.size = Pt(font_size)
    series.data_labels.font.name = H.FONT
    series.data_labels.font.bold = True
    series.data_labels.font.color.rgb = H.INK
    for i, pt in enumerate(series.points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i]
        pt.format.line.fill.background()
        pt.data_label.text_frame.text = value_fmt.format(values[i])

    _style_val_axis_hidden(chart)
    _no_gridlines(chart)
    cax = chart.category_axis
    cax.tick_labels.font.size = Pt(cat_font_size)
    cax.tick_labels.font.name = H.FONT
    cax.tick_labels.font.color.rgb = H.INK2
    cax.format.line.color.rgb = H.RULE
    return gframe


# ------------------------------------------------------------------ vertical bar
def vbar(slide, x, y, w, h, categories, values, colors, value_fmt="{:.0f}", font_size=8.5):
    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("v", values)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                     Inches(w), Inches(h), cd)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 35
    series = chart.series[0]
    series.has_data_labels = True
    series.data_labels.font.size = Pt(font_size)
    series.data_labels.font.name = H.FONT
    series.data_labels.font.bold = True
    series.data_labels.font.color.rgb = H.INK
    for i, pt in enumerate(series.points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i]
        pt.format.line.fill.background()
        pt.data_label.text_frame.text = value_fmt.format(values[i])
    _style_val_axis_hidden(chart)
    _no_gridlines(chart)
    cax = chart.category_axis
    cax.tick_labels.font.size = Pt(font_size - 1)
    cax.tick_labels.font.name = H.FONT
    cax.tick_labels.font.color.rgb = H.INK2
    cax.format.line.color.rgb = H.RULE
    return gframe


# ------------------------------------------------------------------ single-row stacked bar
def stacked_hbar_single(slide, x, y, w, h, seg_names, seg_values, seg_colors,
                         value_fmt="{:.0f}", font_size=10):
    cd = CategoryChartData()
    cd.categories = [""]
    for name, val in zip(seg_names, seg_values):
        cd.add_series(name, (val,))
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(x), Inches(y),
                                     Inches(w), Inches(h), cd)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 20
    plot.overlap = 100
    for i, s in enumerate(chart.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = seg_colors[i]
        s.format.line.color.rgb = H.SURF
        s.format.line.width = Pt(1.5)
        s.has_data_labels = True
        s.data_labels.show_value = True
        s.data_labels.number_format = "0"
        s.data_labels.number_format_is_linked = False
        s.data_labels.font.size = Pt(font_size)
        s.data_labels.font.bold = True
        s.data_labels.font.name = H.FONT
        s.data_labels.font.color.rgb = H.SURF if seg_colors[i] not in (H.DEEMPH,) else H.INK
        try:
            s.points[0].data_label.text_frame.word_wrap = False
        except Exception:
            pass
    _style_val_axis_hidden(chart)
    _no_gridlines(chart)
    chart.category_axis.visible = False
    return gframe


# ------------------------------------------------------------------ small multiples (line, no markers needed but ok)
def small_multiple_line(slide, x, y, w, h, categories, values, color, title, rate_label):
    box, tf = H.add_textbox(slide, x, y, w, 0.15)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(8); r.font.bold = True; r.font.name = H.FONT
    r.font.color.rgb = H.INK if color == H.ACCENT else H.INK2
    box2, tf2 = H.add_textbox(slide, x, y + 0.15, w, 0.13)
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run(); r2.text = rate_label
    r2.font.size = Pt(7); r2.font.bold = True; r2.font.name = H.FONT
    r2.font.color.rgb = color

    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("v", values)
    chart_y = y + 0.30
    chart_h = h - 0.30
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(chart_y),
                                     Inches(w), Inches(chart_h), cd)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    s = chart.series[0]
    s.format.line.color.rgb = color
    s.format.line.width = Pt(1.5)
    _style_val_axis_hidden(chart)
    _no_gridlines(chart)
    cax = chart.category_axis
    cax.tick_labels.font.size = Pt(6.5)
    cax.tick_labels.font.name = H.FONT
    cax.tick_labels.font.color.rgb = H.MUTED
    cax.format.line.color.rgb = H.GRID
    return gframe


# ------------------------------------------------------------------ confidence band (native shapes, not a "chart" per se)
def confidence_band(slide, x, y, w, h, p10, p50, p90, vmin, vmax, label="Paid FTE"):
    track_y = y + h * 0.35
    track_h = h * 0.22
    def xpos(v):
        return x + (v - vmin) / (vmax - vmin) * w
    # full track (light)
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(track_y),
                                 Inches(w), Inches(track_h))
    bg.fill.solid(); bg.fill.fore_color.rgb = H.WASH; bg.line.fill.background()
    bg.shadow.inherit = False
    # band px10-p90 slightly darker
    bx = xpos(p10)
    bw = xpos(p90) - xpos(p10)
    band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(track_y),
                                   Inches(bw), Inches(track_h))
    band.fill.solid(); band.fill.fore_color.rgb = H.ACCENT
    band.fill.fore_color.brightness = 0.55
    band.line.fill.background()
    band.shadow.inherit = False
    # median marker
    mx = xpos(p50)
    marker = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(mx - 0.025),
                                     Inches(track_y - 0.05), Inches(0.05), Inches(track_h + 0.10))
    marker.fill.solid(); marker.fill.fore_color.rgb = H.ACCENT; marker.line.fill.background()
    marker.shadow.inherit = False
    # labels
    def label_at(v, text, y_off, bold=True, size=9, color=None):
        box, tf = H.add_textbox(slide, xpos(v) - 0.4, track_y + y_off, 0.8, 0.16)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = H.FONT
        r.font.color.rgb = color or H.INK2
    label_at(p50, str(p50), -0.28, size=11, color=H.INK)
    label_at(p10, f"P10\n{p10}", track_h + 0.06, size=7, color=H.MUTED)
    label_at(p90, f"P90\n{p90}", track_h + 0.06, size=7, color=H.MUTED)
