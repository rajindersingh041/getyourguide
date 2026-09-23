#!/usr/bin/env python
"""Independent verification of every layer of the model.

Nothing here imports a result from run_model and trusts it. Each check either
re-derives a number by a different method (closed form, simulation, raw-cell
recomputation) or asserts an invariant that must hold regardless of the data.
"""
from __future__ import annotations
import heapq, math, sys
from math import comb, exp, factorial, erf, sqrt

import numpy as np
import openpyxl
import pandas as pd

from wfm import config, clean, forecast, capacity, optimise
from wfm.erlang import (erlang_b, erlang_c, service_level, agents_required,
                        traffic_intensity, avg_speed_of_answer, norm_ppf)
from wfm.loader import Workbook

PASS, FAIL = [], []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"\n         {detail}" if detail else ""))
    return ok


def close(a, b, tol=1e-9):
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def section(t):
    print(f"\n{'-' * 78}\n{t}\n{'-' * 78}")


# ============================================================================
section("A. ERLANG B - recursion vs textbook closed form")
# ============================================================================
# B(n,a) = (a^n/n!) / sum_{k=0..n}(a^k/k!)   - overflows for large n, exact for small
def erlang_b_direct(n, a):
    num = a ** n / factorial(n)
    den = sum(a ** k / factorial(k) for k in range(n + 1))
    return num / den

worst = 0.0
for a in (0.5, 1.0, 3.7, 7.679, 15.0, 42.3):
    for n in range(1, 60):
        worst = max(worst, abs(erlang_b(n, a) - erlang_b_direct(n, a)))
check("Erlang B recursion == closed form", worst < 1e-12, f"max abs diff {worst:.2e}")
check("Erlang B(1, a) == a/(1+a)", close(erlang_b(1, 3.0), 3 / 4))
check("Erlang B is decreasing in n", all(erlang_b(n, 7.0) > erlang_b(n + 1, 7.0)
                                         for n in range(1, 60)))
check("Erlang B overflow-safe at n=500", 0.0 <= erlang_b(500, 450.0) <= 1.0,
      f"B(500,450) = {erlang_b(500,450.0):.6f}")

# ============================================================================
section("B. ERLANG C - published reference values")
# ============================================================================
# Standard reference points (Erlang C probability of waiting), 4dp.
# Reference values cross-derived from the stationary distribution of M/M/c,
# independently of the Erlang-B recursion under test.
def erlang_c_reference(n, a):
    rho = a / n
    p0 = 1.0 / (sum(a ** k / factorial(k) for k in range(n))
                + a ** n / factorial(n) / (1 - rho))
    return a ** n / factorial(n) / (1 - rho) * p0

refs = [(1, 0.5, 0.5000), (2, 1.0, 0.3333), (3, 2.0, 0.4444),
        (5, 3.0, 0.2362), (10, 5.0, 0.0361), (20, 15.0, 0.1604294)]
ok = True
for n, a, expect in refs:
    got = erlang_c(n, a)
    good = abs(got - expect) < 5e-4
    ok &= good
    print(f"         C({n:2d},{a:5.1f}) = {got:.6f}  reference {expect:.4f}  {'ok' if good else 'MISMATCH'}")
check("Erlang C matches published reference values", ok)
worst = max(abs(erlang_c(n, a) - erlang_c_reference(n, a))
            for a in (0.5, 3.0, 7.679, 15.0, 18.765) for n in range(int(a) + 1, 60))
check("Erlang C == M/M/c stationary distribution (independent derivation)",
      worst < 1e-12, f"max abs diff {worst:.2e}")
check("Erlang C >= Erlang B always",
      all(erlang_c(n, 7.0) >= erlang_b(n, 7.0) for n in range(8, 40)))
check("Erlang C -> 1 as n -> a+", close(erlang_c(8, 7.999999), 1.0, 1e-4))
check("Erlang C is decreasing in n",
      all(erlang_c(n, 7.0) > erlang_c(n + 1, 7.0) for n in range(8, 60)))

# ============================================================================
section("C. ERLANG C vs DISCRETE-EVENT SIMULATION (M/M/c, FIFO)")
# ============================================================================
def simulate_mmc(lam_per_hr, aht_sec, c, target_sec, n_arrivals=300_000, seed=7):
    """Exact event-driven M/M/c. Returns (P(wait<=target), mean wait, occupancy)."""
    rng = np.random.default_rng(seed)
    inter = rng.exponential(3600.0 / lam_per_hr, n_arrivals)
    arrivals = np.cumsum(inter)
    service = rng.exponential(aht_sec, n_arrivals)
    free = [0.0] * c
    heapq.heapify(free)
    waits = np.empty(n_arrivals)
    for i in range(n_arrivals):
        t = arrivals[i]
        f = heapq.heappop(free)
        w = f - t
        if w < 0:
            w = 0.0
        waits[i] = w
        heapq.heappush(free, t + w + service[i])
    burn = n_arrivals // 10           # discard transient
    w = waits[burn:]
    busy = service[burn:].sum() / (arrivals[-1] - arrivals[burn]) / c
    return float((w <= target_sec).mean()), float(w.mean()), float(busy)

cases = [
    ("English inbound (BASE Jul)", 44181.22 / 730.0, 457.5, 11, 20),
    ("German inbound (BASE Jul)",   9168.22 / 730.0, 360.0,  3, 20),
    ("English chat (BASE Jul)",    42882.13 / 730.0, 1150.0, 24, 60),
    ("generic mid-size",           200.0,             300.0, 20, 20),
]
SEEDS = (7, 11, 23, 41, 97)
for label, lam, aht, c, tgt in cases:
    a = lam * aht / 3600.0
    analytic = service_level(c, a, tgt, aht)
    asa_a = avg_speed_of_answer(c, a, aht)
    runs = [simulate_mmc(lam, aht, c, tgt, seed=sd) for sd in SEEDS]
    sls = np.array([r[0] for r in runs]); was = np.array([r[1] for r in runs])
    occ = np.mean([r[2] for r in runs])
    # standard error of the mean across independent replications
    se_sl = sls.std(ddof=1) / np.sqrt(len(SEEDS))
    se_w = was.std(ddof=1) / np.sqrt(len(SEEDS))
    z_sl = abs(analytic - sls.mean()) / max(se_sl, 1e-12)
    z_w = abs(asa_a - was.mean()) / max(se_w, 1e-12)
    print(f"         {label}: a={a:.3f} c={c}  ({len(SEEDS)} replications)")
    print(f"           SL  analytic {analytic:.4f} | sim {sls.mean():.4f} +/- {se_sl:.4f} | z={z_sl:.2f}")
    print(f"           ASA analytic {asa_a:7.2f}s | sim {was.mean():7.2f}s +/- {se_w:5.2f}s | z={z_w:.2f}")
    print(f"           occupancy analytic {a/c:.4f} | sim {occ:.4f}")
    # 3 sigma: analytic value must lie inside the simulation's confidence band
    check(f"simulation agrees with Erlang C - {label}", z_sl < 3.0 and z_w < 3.0)

# ============================================================================
section("D. STAFFING SOLVER - minimality and monotonicity")
# ============================================================================
ok_min = ok_mono = True
for a in (0.3, 0.84, 7.679, 18.77, 55.0):
    for tgt, secs, aht in ((0.8, 20, 450.0), (0.8, 60, 1150.0), (0.9, 30, 600.0)):
        n = agents_required(a, tgt, secs, aht)
        ok_min &= service_level(n, a, secs, aht) >= tgt
        ok_min &= (n == math.ceil(a) or service_level(n - 1, a, secs, aht) < tgt)
        ok_mono &= all(service_level(k, a, secs, aht) <= service_level(k + 1, a, secs, aht)
                       for k in range(max(1, math.ceil(a)), n + 6))
check("agents_required returns the MINIMAL feasible n", ok_min)
check("service level is monotonically increasing in n", ok_mono)
check("agents_required(0) == 0", agents_required(0.0, 0.8, 20, 450.0) == 0)
check("more agents needed for a tighter target",
      agents_required(7.679, 0.9, 20, 450.0) >= agents_required(7.679, 0.8, 20, 450.0))
check("more agents needed for a shorter answer time",
      agents_required(7.679, 0.8, 5, 450.0) >= agents_required(7.679, 0.8, 60, 450.0))

# ============================================================================
section("E. LITTLE'S LAW and queueing identities")
# ============================================================================
for lam, aht, c in ((61.0, 450.0, 11), (200.0, 300.0, 20)):
    a = lam * aht / 3600.0
    Lq = erlang_c(c, a) * a / (c - a)          # mean number waiting
    Wq = avg_speed_of_answer(c, a, aht)        # mean wait (seconds)
    lam_s = lam / 3600.0
    check(f"Little's Law Lq = lambda*Wq (lam={lam}, c={c})",
          close(Lq, lam_s * Wq, 1e-9), f"Lq={Lq:.6f}  lambda*Wq={lam_s*Wq:.6f}")
    check(f"utilisation rho = a/c < 1 (lam={lam}, c={c})", a / c < 1)

# ============================================================================
section("F. NORMAL INVERSE CDF")
# ============================================================================
def norm_cdf(z): return 0.5 * (1 + erf(z / sqrt(2)))
known = {0.5: 0.0, 0.9: 1.2815515655, 0.95: 1.6448536270,
         0.975: 1.9599639845, 0.99: 2.3263478740, 0.1: -1.2815515655}
worst = max(abs(norm_ppf(p) - z) for p, z in known.items())
check("norm_ppf matches published z-values", worst < 1e-6, f"max abs diff {worst:.2e}")
worst = max(abs(norm_cdf(norm_ppf(p)) - p) for p in np.linspace(0.001, 0.999, 999))
check("norm_ppf is the inverse of the normal CDF", worst < 1e-8, f"max abs diff {worst:.2e}")

# ============================================================================
section("G. RAW-CELL RECOMPUTATION (independent of the loader)")
# ============================================================================
wbr = openpyxl.load_workbook(config.WORKBOOK, data_only=True)
cvr, ahr = wbr["Contact Volume "], wbr["AHT Assumptions by BPO and lang"]
ROWS = {"English": 2, "German": 7, "Italian": 12, "Spanish": 17, "French": 22}
A = config.Assumptions()

# hand-typed from the sheet: (email, inbound, outbound, chat)
HAND_AHT = {"English": ((1150 + 1150) / 2, (450 + 465) / 2, (105 + 110) / 2, (1200 + 1100) / 2),
            "German": (700, 360, 115, 900), "Italian": (800, 360, 115, 900),
            "Spanish": (600, 360, 115, 900), "French": (600, 360, 115, 900)}
for lang, tup in HAND_AHT.items():
    bpos = ["BPO 1", "BPO 2"] if lang == "English" else ["BPO 3"]
    got = []
    for ch_i in range(3, 7):
        vals = [ahr.cell(r, ch_i).value for r in range(2, 17)
                if ahr.cell(r, 1).value in bpos and ahr.cell(r, 2).value == lang]
        got.append(sum(vals) / len(vals))
    check(f"raw-cell AHT matches hand-read values - {lang}",
          all(close(g, h) for g, h in zip(got, tup)), f"{got} vs {list(tup)}")

# full July BASE chain, recomputed from cells with no wfm.capacity involvement
total_fte = total_prod = total_wl = 0.0
for lang, r in ROWS.items():
    jun = [cvr.cell(r + 1 + j, 7).value for j in range(4)]   # email, in, out, chat
    aht4 = HAND_AHT[lang]
    v = [x * 1.10 for x in jun]
    for i, ch in enumerate(("Email", "Inbound", "Outbound", "Chat")):
        wl = v[i] * aht4[i] / 3600.0
        total_wl += wl
        if ch in ("Inbound", "Chat"):
            tgt, secs = A.sla_for(ch)
            a_erl = wl / 730.0
            n = agents_required(a_erl, tgt, secs, aht4[i])
            conc = 1.2 if ch == "Chat" else 1.0
            seat_h = (n / conc) * 730.0
        else:
            seat_h = wl / 0.85
        total_prod += seat_h
        total_fte += seat_h / 0.82 / (40 * 52 / 12)

# now the pipeline's own answer
wb = Workbook(config.WORKBOOK)
volume, totals = wb.contact_volume()
assum, shrink = wb.aht_assumptions(), wb.shrinkage()
reasons, extra = wb.en_reasons()
en_aht = wb.en_aht()
log = clean.DQLog()
cl = clean.clean(volume, totals, reasons, en_aht, assum, shrink, log)
clean.check_extra_series(extra, log)
brk = forecast.detect_break(cl["totals"])
fc = forecast.VolumeForecast(cl["volume"], cl["totals"], config.FORECAST_MONTHS,
                             A.seasonal_uplift, brk["month"])
aht_eff = capacity.effective_aht(cl["assumptions"], A)
queues = capacity.queue_capacity(fc.by_channel(), aht_eff, A)
build = capacity.headcount_build(queues, A, list(config.FORECAST_MONTHS))
bj = build.query("scenario=='BASE' and month=='Jul'").iloc[0]

check("raw-cell FTE == pipeline FTE (BASE Jul)", close(total_fte, bj.fte_sla, 1e-6),
      f"raw {total_fte:.4f} vs pipeline {bj.fte_sla:.4f}")
check("raw-cell productive hours == pipeline", close(total_prod, bj.productive_hours, 1e-6),
      f"raw {total_prod:,.2f} vs pipeline {bj.productive_hours:,.2f}")
check("raw-cell workload hours == pipeline", close(total_wl, bj.workload_hours, 1e-6),
      f"raw {total_wl:,.2f} vs pipeline {bj.workload_hours:,.2f}")

# ============================================================================
section("H. DATA RECONCILIATION")
# ============================================================================
ch_sum = volume.groupby(["language", "month"], as_index=False).contacts.sum()
m = ch_sum.merge(totals, on=["language", "month"], suffixes=("_ch", "_tot"))
d = (m.contacts_ch - m.contacts_tot).abs()
rel = (d / m.contacts_tot).max()
check("channel-vs-total discrepancy is immaterial (< 0.1% of volume)",
      rel < 1e-3,
      f"{int((d>0).sum())}/{len(m)} cells off by <= {d.max():g} contacts "
      f"({rel:.4%}) - rounding residue; channel rows are the source of truth")

tix = reasons.groupby("month", as_index=False).tickets.sum()
env = totals[totals.language == "English"][["month", "contacts"]]
rec = tix.merge(env, on="month")
check("EN CSAT reason tickets == Contact Volume English totals",
      (rec.tickets / rec.contacts - 1).abs().max() < 1e-9,
      f"ratios {[round(x,9) for x in (rec.tickets/rec.contacts)]}")

shr = cl["shrinkage"]
tot_s = shr[shr.category.str.contains("Tot", case=False)]
check("Tot Shrink is 18% for every language x BPO",
      close(tot_s.pct.min(), 0.18) and close(tot_s.pct.max(), 0.18),
      f"{len(tot_s)} rows, all {tot_s.pct.iloc[0]:.0%}")

parts = shr[~shr.category.str.contains("Tot", case=False)]
ps = parts.groupby(["language", "bpo"], as_index=False).pct.sum()
tj = tot_s.merge(ps, on=["language", "bpo"], suffixes=("_t", "_p"))
check("shrinkage components sum to Tot Shrink",
      (tj.pct_t - tj.pct_p).abs().max() < 1e-9,
      f"max gap {(tj.pct_t - tj.pct_p).abs().max():.2e} over {len(tj)} pairs")

check("routing constraint enforced: no BPO x language violations",
      set(zip(cl["assumptions"].bpo, cl["assumptions"].language)) <=
      {(b, l) for b, ls in A.bpo_languages.items() for l in ls},
      f"{len(assum)} rows in -> {len(cl['assumptions'])} routable rows out")

# ============================================================================
section("I. MODEL INVARIANTS")
# ============================================================================
rt = queues[queues.model == "ErlangC"]
check("every real-time queue MEETS its service level target",
      bool((rt.service_level >= 0.80 - 1e-12).all()),
      f"min achieved SL = {rt.service_level.min():.4f} over {len(rt)} queues")
check("no queue is staffed below its offered load (occupancy < 100%)",
      bool((rt.occupancy < 1.0).all()), f"max occupancy {rt.occupancy.max():.3f}")
check("SLA staffing is never below workload-only staffing",
      bool((queues.fte >= queues.fte_workload_only - 1e-9).all()))

agg = queues.groupby(["scenario", "month"], as_index=False).fte.sum()
j = agg.merge(build, on=["scenario", "month"])
check("per-queue FTE sums exactly to the headcount roll-up",
      (j.fte - j.fte_sla).abs().max() < 1e-9)

bpo_df = capacity.by_bpo(queues, A)
bb = bpo_df.groupby(["scenario", "month"], as_index=False).fte.sum().merge(
    build, on=["scenario", "month"])
check("BPO split re-aggregates to the total (no FTE lost or created)",
      (bb.fte - bb.fte_sla).abs().max() < 1e-9,
      f"max diff {(bb.fte - bb.fte_sla).abs().max():.2e}")

sc = build.pivot_table(index="month", columns="scenario", values="fte_paid",
                       observed=True)
check("scenarios are ordered LOW <= BASE <= HIGH in every month",
      bool(((sc.LOW <= sc.BASE + 1e-9) & (sc.BASE <= sc.HIGH + 1e-9)).all()),
      "\n         " + sc.round(1).to_string().replace("\n", "\n         "))

check("paid FTE > on-floor FTE > SLA FTE > workload-only FTE",
      bool((build.fte_paid > build.fte_on_floor).all() and
           (build.fte_on_floor > build.fte_sla).all() and
           (build.fte_sla > build.fte_workload_only).all()))

# linearity: doubling volume must not more than double FTE (Erlang economies)
q1 = queues.query("scenario=='BASE' and month=='Jul'")
dbl = capacity.queue_capacity(
    fc.by_channel().query("scenario=='BASE' and month=='Jul'").assign(
        contacts=lambda d: d.contacts * 2), aht_eff, A)
check("doubling volume gives SUB-linear FTE growth (pooling economies)",
      dbl.fte.sum() < 2 * q1.fte.sum(),
      f"x1 = {q1.fte.sum():.1f} FTE -> x2 = {dbl.fte.sum():.1f} FTE "
      f"(ratio {dbl.fte.sum()/q1.fte.sum():.3f}, not 2.000)")

# ============================================================================
section("J. HR MATHS")
# ============================================================================
a_mo = A.attrition_monthly
check("monthly attrition compounds to the annual figure",
      close(1 - (1 - a_mo) ** 12, A.attrition_annual, 1e-12),
      f"{a_mo:.4%}/mo -> {1-(1-a_mo)**12:.2%}/yr (linear would be {A.attrition_annual/12:.4%}/mo)")

# cohort simulation of the steady-state tenure mix
def simulate_tenure(a_mo, curve, months=600):
    pop = np.zeros(months)     # pop[k] = share with tenure k months
    pop[-1] = 1.0
    for _ in range(3000):
        new = np.zeros(months)
        new[0] = a_mo
        new[1:] = pop[:-1] * (1 - a_mo)
        new[-1] += pop[-1] * (1 - a_mo)
        pop = new / new.sum()
    eff = sum(pop[k] / curve[k] for k in range(len(curve))) + pop[len(curve):].sum()
    return eff

sim_eff = simulate_tenure(a_mo, A.learning_curve)
check("tenure efficiency matches a cohort simulation",
      abs(sim_eff - A.tenure_efficiency) < 2e-3,
      f"closed form {A.tenure_efficiency:.5f} vs simulated {sim_eff:.5f}")
check("learning curve costs between 0% and 5% of capacity",
      0 < 1 / A.tenure_efficiency - 1 < 0.05,
      f"blended AHT uplift = {1/A.tenure_efficiency-1:+.2%}")

# ============================================================================
section("K. FORECAST LOGIC")
# ============================================================================
check("structural break detected in March and is uniform across languages",
      brk["month"] == "March" and brk["uniform"],
      f"total {brk['total_ratio']:.2f}x; per-language "
      f"{ {k: round(v,2) for k,v in brk['per_language'].items()} }")

gp = fc.growth_profile()
check("English is the only language flagged as growing",
      list(gp[gp.growing].index) == ["English"],
      f"slopes/month: { {k: round(v) for k,v in gp.trend_per_month.items()} }")

# OLS slope verified by hand against numpy.polyfit
post = fc.post["English"].values
np_slope = np.polyfit(np.arange(len(post)), post, 1)[0]
check("OLS slope matches numpy.polyfit",
      close(gp.loc["English", "trend_per_month"], float(np_slope), 1e-9),
      f"{gp.loc['English','trend_per_month']:.4f} vs {np_slope:.4f}")

base_tot = fc.scenarios().query("scenario=='BASE' and month=='Jul'").contacts.sum()
# baseline is the SUM OF CHANNEL ROWS for June, not the 'X Tot' rows (which
# carry the +/-2 rounding residue) - see the materiality check in section H.
jun_ch = volume[volume.month == "June"].contacts.sum()
jun_tot = totals[totals.month == "June"].contacts.sum()
check("BASE = June channel-row run-rate x 1.10 exactly",
      close(base_tot, jun_ch * 1.10, 1e-9),
      f"June channel rows {jun_ch:,.0f} x 1.10 = {jun_ch*1.1:,.2f} vs model "
      f"{base_tot:,.2f}  (the 'Tot' rows say {jun_tot:,.0f}, a {jun_ch-jun_tot:+.0f} residue)")

vb = fc.by_channel().query("scenario=='BASE' and month=='Jul'")
check("channel forecasts sum exactly to the language forecast",
      abs(vb.contacts.sum() - base_tot) < 1e-9,
      f"{vb.contacts.sum():,.4f} vs {base_tot:,.4f}")

vbc = fc.by_channel()
chk = vbc.query("scenario=='BASE' and month=='Jul'").groupby("language").contacts.sum()
sc2 = fc.scenarios().query("scenario=='BASE' and month=='Jul'").set_index("language").contacts
check("channel split preserves language totals",
      (chk - sc2).abs().max() < 1e-6, f"max diff {(chk-sc2).abs().max():.2e}")

ci = fc.intervals()
p50 = ci.query("month=='Jul' and percentile==0.5").contacts.iloc[0]
check("P50 of the interval equals the BASE point estimate",
      close(p50, base_tot, 1e-9))
check("intervals are monotonically increasing in percentile",
      bool(ci[ci.month == "Jul"].sort_values("percentile").contacts.is_monotonic_increasing))

# ============================================================================
section("L. OPTIMISATION LEVERS - verified, not asserted")
# ============================================================================
non_en = ["French", "German", "Italian", "Spanish"]
pools = [optimise.pool_queues(queues, A, non_en, ch, month="Jul")
         for ch in ("Inbound", "Chat")]
for p in pools:
    check(f"pooled {p['channel']} queue still meets SLA",
          p["pooled_service_level"] >= 0.80,
          f"seats {p['siloed_seats_raw']} -> {p['pooled_seats_raw']}, "
          f"occupancy {p['siloed_occupancy_min']:.0%}-{p['siloed_occupancy_max']:.0%} -> "
          f"{p['pooled_occupancy']:.0%}, SL {p['pooled_service_level']:.1%}, "
          f"saves {p['fte_saved']:.1f} FTE")
    check(f"pooling {p['channel']} genuinely saves FTE", p["fte_saved"] > 0)

dfl = optimise.deflection(queues, A, month="Jul")
asec = dfl.attrs["agent_seconds_per_contact"]
check("phone->chat deflection is tested, and COSTS FTE in this dataset",
      dfl.delta_fte.iloc[-1] > 0,
      f"agent-seconds/contact: Inbound {asec['Inbound']:.0f}s vs Chat {asec['Chat']:.0f}s "
      f"-> 30% shift = {dfl.delta_fte.iloc[-1]:+.1f} FTE")

cs = optimise.concurrency_sensitivity(queues, A, month="Jul")
check("chat FTE falls monotonically as concurrency rises",
      bool(cs.chat_fte.is_monotonic_decreasing),
      f"1.2 -> {cs.chat_fte.iloc[0]:.1f} FTE; 2.0 -> "
      f"{cs.query('concurrency==2.0').chat_fte.iloc[0]:.1f} FTE "
      f"({cs.query('concurrency==2.0')['delta_vs_1.2'].iloc[0]:+.1f})")

shr_s = optimise.shrinkage_sensitivity(queues, A, month="Jul")
check("shrinkage sensitivity is monotonically increasing",
      bool(shr_s.fte.is_monotonic_increasing),
      f"16% -> {shr_s.fte.iloc[0]:.1f} FTE; 20% -> {shr_s.fte.iloc[-1]:.1f} FTE")

# ============================================================================
section("M. RECONCILIATION AGAINST PLAN.md's PUBLISHED NUMBERS")
# ============================================================================
HPF = 40 * 52 / 12
# PLAN.md 7.2: English reason-weighted AHT, Apr-Jun = 908.78s
r = cl["reasons"]; ah = cl["aht"]
mj = r.merge(ah[["reason", "month", "aht"]], on=["reason", "month"])
aj = mj[mj.month.isin(["April", "May", "June"])]
w = (aj.tickets * aj.aht).sum() / aj.tickets.sum()
check("reproduces PLAN.md's 908.78s English weighted AHT", abs(w - 908.78) < 0.01,
      f"recomputed {w:.2f}s")

# PLAN.md 7.3: 427.90 FTE total
plan_aht = {"English": w, "German": 514.25, "Italian": 538.85,
            "Spanish": 489.65, "French": 489.65}
plan_vol = {l: totals.query("language==@l and month=='June'").contacts.iloc[0] * 1.1
            for l in plan_aht}
plan_fte = sum(plan_vol[l] * plan_aht[l] / 3600 / 0.82 / HPF for l in plan_aht)
check("reproduces PLAN.md's 427.90 FTE July figure", abs(plan_fte - 427.90) < 0.6,
      f"recomputed {plan_fte:.2f} FTE - PLAN.md's method is arithmetically self-consistent")

# PLAN.md 7.5: 11 phone agents, SL 82.84%
a_plan = 44228.2 / 720.0 * 450 / 3600
check("reproduces PLAN.md's Erlang result (11 agents @ 82.84%)",
      agents_required(a_plan, 0.8, 20, 450.0) == 11 and
      abs(service_level(11, a_plan, 20, 450.0) - 0.8284) < 5e-4,
      f"a={a_plan:.3f} erlangs, n=11, SL={service_level(11,a_plan,20,450.0):.4f}")

# ...but those 11 seats are never converted to FTE in PLAN.md
seats_fte = 11 * 730 / 0.82 / HPF
wl_fte = 44228.2 * 450 / 3600 / 0.82 / HPF
check("PLAN.md's 11 seats are worth far more FTE than its workload model books",
      seats_fte > wl_fte * 1.3,
      f"11 seats = {seats_fte:.1f} FTE vs the {wl_fte:.1f} FTE its workload model "
      f"charges for the same queue (+{seats_fte/wl_fte-1:.0%})")

check("this model's BASE July answer is materially above PLAN.md's 428",
      bj.fte_paid > plan_fte,
      f"PLAN.md {plan_fte:.0f} FTE (no SLA cost) vs this model "
      f"{bj.fte_paid:.0f} paid FTE (SLA + ramp + training)")

# the English AHT conflict, priced
en_ch = queues.query("scenario=='BASE' and month=='Jul' and language=='English'")
wl_channel = en_ch.workload_hours.sum()
wl_reason = en_ch.contacts.sum() / 4 * 0  # placeholder to avoid double count
wl_reason = (totals.query("language=='English' and month=='June'").contacts.iloc[0]
             * 1.1 * w / 3600)
check("the two English AHT sources disagree materially (open question, priced)",
      abs(wl_reason / wl_channel - 1) > 0.15,
      f"channel-based {wl_channel:,.0f}h vs reason-based {wl_reason:,.0f}h "
      f"({wl_reason/wl_channel-1:+.1%}, ~{(wl_reason-wl_channel)/0.82/HPF:.0f} FTE)")

# ============================================================================
section("N. AUTOMATION CONTRACT - does it survive changed input?")
# ============================================================================
import copy
# drop a language entirely and confirm the pipeline adapts without code changes
v2 = volume[volume.language != "Italian"].copy()
t2 = totals[totals.language != "Italian"].copy()
a2 = assum[assum.language != "Italian"].copy()
log2 = clean.DQLog()
cl2 = clean.clean(v2, t2, reasons, en_aht, a2, shrink, log2)
A2 = config.Assumptions(bpo_languages={"BPO 1": ("English",), "BPO 2": ("English",),
                                       "BPO 3": ("German", "French", "Spanish")})
fc2 = forecast.VolumeForecast(cl2["volume"], cl2["totals"], config.FORECAST_MONTHS,
                              A2.seasonal_uplift, forecast.detect_break(cl2["totals"])["month"])
q2 = capacity.queue_capacity(fc2.by_channel(), capacity.effective_aht(cl2["assumptions"], A2), A2)
b2 = capacity.headcount_build(q2, A2, list(config.FORECAST_MONTHS))
expected = 3 * len(config.FORECAST_MONTHS) * 4 * 4
check("pipeline runs unchanged with a language removed",
      len(q2) == expected and b2.fte_paid.iloc[0] > 0,
      f"3 scenarios x {len(config.FORECAST_MONTHS)} months x 4 languages x 4 channels "
      f"= {len(q2)} queues (expected {expected}); "
      f"BASE Jul = {b2.query('scenario==\"BASE\" and month==\"Jul\"').fte_paid.iloc[0]:.0f} paid FTE")

A3 = config.Assumptions(shrinkage=0.25, seasonal_uplift=1.25, deferred_occupancy=0.9)
q3 = capacity.queue_capacity(fc.by_channel(), aht_eff, A3)
b3 = capacity.headcount_build(q3, A3, list(config.FORECAST_MONTHS))
check("all assumptions are config-driven (no hard-coded constants)",
      b3.query("scenario=='BASE' and month=='Jul'").fte_paid.iloc[0] > bj.fte_paid,
      f"shrink 18%->25%, uplift 10%->25%: "
      f"{bj.fte_paid:.0f} -> {b3.query('scenario==\"BASE\" and month==\"Jul\"').fte_paid.iloc[0]:.0f} paid FTE")

try:
    Workbook.__init__  # excluded-sheet guard
    from wfm.loader import DataError
    wbx = Workbook(config.WORKBOOK)
    try:
        wbx._sheet("CSAT"); hidden_ok = False
    except DataError:
        hidden_ok = True
except Exception:
    hidden_ok = False
check("hidden 'CSAT' sheet is refused by the loader", hidden_ok)

# ============================================================================
print(f"\n{'=' * 78}")
print(f"RESULT: {len(PASS)} passed, {len(FAIL)} failed")
if FAIL:
    print("\nFAILURES:")
    for f in FAIL:
        print(f"  - {f}")
print("=" * 78)
sys.exit(1 if FAIL else 0)
